#!/usr/bin/env python3
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

SKILL_ROOT = Path(
    "/Users/haozhanyi/.codex/plugins/cache/openai-primary-runtime/"
    "documents/26.601.10930/skills/documents"
)
sys.path.insert(0, str(SKILL_ROOT / "scripts"))
from table_geometry import apply_table_geometry  # noqa: E402

OUTPUT = Path(
    "/Users/haozhanyi/Documents/claude/release/manual/"
    "AI工具部署助手-客户交付与使用手册.docx"
)

NAVY = "142033"
BLUE = "2E74B5"
ORANGE = "D65D36"
LIGHT_BLUE = "E8EEF5"
LIGHT_ORANGE = "F9E8E0"
LIGHT_GREEN = "E7F3EC"
LIGHT_GRAY = "F2F4F7"
MUTED = "667085"
WHITE = "FFFFFF"
BLACK = "1F2937"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = "D7DEE8") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        border = borders.find(qn(tag))
        if border is None:
            border = OxmlElement(tag)
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), color)


def set_run_font(
    run,
    *,
    size: float = 11,
    bold: bool | None = None,
    color: str = BLACK,
    italic: bool | None = None,
) -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "PingFang SC")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    specs = {
        "Title": (30, NAVY, 0, 10),
        "Subtitle": (14, MUTED, 0, 18),
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, "1F4D78", 10, 5),
    }
    for name, (size, color, before, after) in specs.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = name.startswith("Heading")


def add_numbering(doc: Document) -> tuple[int, int]:
    numbering = doc.part.numbering_part.element

    def abstract_num(abstract_id: int, fmt: str, text: str, font: str | None) -> None:
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        suffix = OxmlElement("w:suff")
        suffix.set(qn("w:val"), "tab")
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "540")
        ind.set(qn("w:hanging"), "270")
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "80")
        spacing.set(qn("w:line"), "300")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.extend([tabs, ind, spacing])
        lvl.extend([start, num_fmt, lvl_text, suffix, p_pr])
        if font:
            r_pr = OxmlElement("w:rPr")
            fonts = OxmlElement("w:rFonts")
            fonts.set(qn("w:ascii"), font)
            fonts.set(qn("w:hAnsi"), font)
            r_pr.append(fonts)
            lvl.append(r_pr)
        abstract.append(lvl)
        numbering.append(abstract)

    def num(num_id: int, abstract_id: int) -> None:
        element = OxmlElement("w:num")
        element.set(qn("w:numId"), str(num_id))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), str(abstract_id))
        element.append(abstract_ref)
        numbering.append(element)

    abstract_num(42, "decimal", "%1.", None)
    abstract_num(43, "bullet", "•", "Arial")
    num(42, 42)
    num(43, 43)
    return 42, 43


def numbered_paragraph(doc: Document, text: str, num_id: int) -> None:
    paragraph = doc.add_paragraph()
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])
    p_pr.append(num_pr)
    set_run_font(paragraph.add_run(text))


def add_callout(doc: Document, title: str, body: str, fill: str = LIGHT_ORANGE) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.12)
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "6")
        border.set(qn("w:space"), "8")
        border.set(
            qn("w:color"),
            "D8C2B8" if fill == LIGHT_ORANGE else "C7D8E8",
        )
        borders.append(border)
    p_pr.extend([shading, borders])
    set_run_font(paragraph.add_run(f"{title}\n"), bold=True, color=ORANGE)
    set_run_font(paragraph.add_run(body), size=10.5)


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[int],
    *,
    header_fill: str = LIGHT_BLUE,
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, header_fill)
        set_cell_borders(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER if len(headers) > 2 else WD_ALIGN_PARAGRAPH.LEFT
        )
        paragraph.paragraph_format.space_after = Pt(0)
        set_run_font(paragraph.add_run(header), size=10, bold=True, color=NAVY)
    for row_values in rows:
        row = table.add_row()
        for index, value in enumerate(row_values):
            cell = row.cells[index]
            set_cell_borders(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
                if len(headers) > 2 and index == 0
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            set_run_font(paragraph.add_run(value), size=9.7)
    apply_table_geometry(table, widths)
    after = doc.add_paragraph()
    after.paragraph_format.space_before = Pt(4)
    after.paragraph_format.space_after = Pt(4)


def add_footer(section) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(
        paragraph.add_run("AI 工具部署助手 | 客户交付与使用手册 | 第 "),
        size=8.5,
        color=MUTED,
    )
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)
    set_run_font(paragraph.add_run(" 页"), size=8.5, color=MUTED)


def add_cover(doc: Document) -> None:
    doc.add_paragraph().paragraph_format.space_after = Pt(72)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(
        kicker.add_run("CUSTOMER DELIVERY & OPERATOR GUIDE"),
        size=10,
        bold=True,
        color=ORANGE,
    )
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("AI 工具部署助手")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("客户交付、激活、安装、模型接入与售后操作手册")
    doc.add_paragraph().paragraph_format.space_after = Pt(44)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(meta.add_run("正式 Release 构建 1.0.0\n"), size=12, bold=True, color=NAVY)
    set_run_font(meta.add_run("适用：macOS 12+ / Apple Silicon\n"), size=11, color=MUTED)
    set_run_font(meta.add_run("文档日期：2026 年 6 月 13 日"), size=10, color=MUTED)
    doc.add_paragraph().paragraph_format.space_after = Pt(74)
    warning = doc.add_paragraph()
    warning.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(
        warning.add_run(
            "第三方部署工具，不隶属于 Anthropic、OpenClaw、Nous Research 或 CC Switch。"
        ),
        size=9.5,
        italic=True,
        color=MUTED,
    )
    doc.add_page_break()


def build() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.right_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)
        add_footer(section)
    configure_styles(doc)
    decimal_num, bullet_num = add_numbering(doc)
    add_cover(doc)

    doc.add_heading("先看结论", level=1)
    add_callout(
        doc,
        "你在闲鱼订单中交付哪个文件？",
        "交付 release/customer-delivery/AI工具部署助手-客户交付包-macOS-1.0.0.zip。"
        "激活码必须在订单聊天中单独发送，不能写进压缩包，也不能把后台管理员 Token 发给客户。",
    )
    add_table(
        doc,
        ["项目", "当前正式交付状态"],
        [
            ["主程序", "macOS Apple Silicon Release 版，必须输入有效激活码"],
            ["配套组件", "CC Switch 3.16.2 macOS Universal 官方 DMG，MIT 许可证"],
            ["支持工具", "Claude Code、OpenClaw、Hermes Agent"],
            ["不包含", "AI 账号、API Key、模型额度、代理、翻墙或地区限制绕过"],
            ["暂不销售", "Windows 主程序尚未生成，不能把 CC Switch Windows 包当作主程序"],
        ],
        [2700, 6660],
    )
    add_callout(
        doc,
        "商业发布前必须确认",
        "当前 Release 构建具备激活校验和签名更新包，但还应确认 Apple Developer ID 签名与公证状态。"
        "如果没有完成 Apple 公证，只能标记为内部交付版，不要宣传为“苹果官方认证安装包”。",
        LIGHT_BLUE,
    )

    doc.add_heading("第一部分：卖家发布与交付", level=1)
    doc.add_heading("1. 闲鱼商品应如何描述", level=2)
    numbered_paragraph(
        doc,
        "标题建议：Mac AI 工具部署助手｜Claude Code / OpenClaw / Hermes 一键安装与配置指导。",
        decimal_num,
    )
    numbered_paragraph(
        doc,
        "商品类型写“软件使用授权 + 安装配置教程”，明确是第三方工具，不写“官方版”“官方授权”或“永久官方账号”。",
        decimal_num,
    )
    numbered_paragraph(
        doc,
        "详情页列出支持系统：macOS 12 及以上、Apple 芯片；购买前让客户发送“关于本机”截图确认芯片。",
        decimal_num,
    )
    numbered_paragraph(
        doc,
        "明确客户自备合法模型账号或 API Key，软件不售卖账号、密钥、共享额度或网络代理。",
        decimal_num,
    )
    numbered_paragraph(
        doc,
        "写清售后范围、授权天数、可重置设备次数和退款边界，所有承诺与后台实际设置一致。",
        decimal_num,
    )

    doc.add_heading("2. 可直接使用的商品详情模板", level=2)
    add_callout(
        doc,
        "商品详情示例",
        "本商品为 macOS Apple 芯片适用的第三方 AI 工具部署助手授权。"
        "可用于安装、诊断和初始化 Claude Code、OpenClaw、Hermes Agent，并附 CC Switch 模型配置工具和中文教程。"
        "购买后交付安装包、使用手册和一机一码激活码。客户需自备合法可用的模型账号或 API Key。"
        "本商品不包含任何官方账号、模型额度、代理、翻墙或地区限制绕过服务；各第三方工具版权归原作者所有。",
        LIGHT_BLUE,
    )
    add_table(
        doc,
        ["详情页必须写", "不要写"],
        [
            ["第三方部署助手、支持系统、授权期限、售后范围", "Anthropic 官方版、官方永久授权"],
            ["客户自备账号/API Key，第三方服务可能收费", "内置无限额度、免费永久使用模型"],
            ["CC Switch 为 MIT 开源配套工具", "把 CC Switch 宣传成自研闭源组件"],
            ["当前只交付 macOS Apple 芯片主程序", "支持 Windows 主程序或所有 Mac"],
        ],
        [4680, 4680],
        header_fill=LIGHT_ORANGE,
    )

    doc.add_heading("3. 每笔订单的标准交付流程", level=2)
    steps = [
        "确认客户设备是 Apple 芯片 Mac，并确认系统版本满足要求。",
        "登录管理后台，在“批量生成兑换码”中生成 1 个兑换码；标签填写闲鱼订单号或客户备注。",
        "上传或发送客户交付 ZIP。推荐使用你控制的 OSS 下载地址或合规网盘链接，并设置合理有效期。",
        "在订单聊天中单独发送兑换码，提醒客户不要公开或转发。",
        "让客户先阅读 PDF 手册，再安装主程序并激活。",
        "客户激活成功后，在后台确认设备绑定和客户端版本。",
        "记录交付日期、授权到期日和售后截止日，不保存客户 API Key。",
    ]
    for step in steps:
        numbered_paragraph(doc, step, decimal_num)

    doc.add_heading("4. 管理后台怎么使用", level=2)
    add_table(
        doc,
        ["操作", "方法", "注意事项"],
        [
            ["打开后台", "访问 https://120-55-15-161.sslip.io/admin", "只在可信设备操作"],
            ["管理员登录", "输入服务端保存的 Admin Token", "不要发给客户，不要写入交付包"],
            ["生成兑换码", "数量 1，填写订单标签、到期时间和重置次数", "完整码只显示一次"],
            ["查看状态", "兑换码列表查看绑定设备、到期日和版本", "不记录客户 API Key"],
            ["客户换机", "使用“重置绑定”后让客户重新激活", "按售后政策使用次数"],
            ["停止授权", "停用兑换码", "停用后续期会被拒绝"],
        ],
        [1600, 4300, 3460],
    )
    add_callout(
        doc,
        "管理员安全红线",
        "不得把 .secrets、管理员 Token、许可证私钥、清单签名私钥、数据库备份或服务器 SSH 凭据放入客户 ZIP。"
        "完整兑换码只通过订单聊天单独交付。",
    )

    doc.add_heading("第二部分：客户安装与激活", level=1)
    doc.add_heading("5. 安装前检查", level=2)
    add_table(
        doc,
        ["检查项", "合格标准"],
        [
            ["电脑", "Apple Silicon Mac（M1、M2、M3、M4 或后续 Apple 芯片）"],
            ["系统", "macOS 12 Monterey 或更高版本"],
            ["网络", "可访问激活服务、安装源和所选模型服务"],
            ["磁盘", "建议至少预留 5 GB"],
            ["模型服务", "客户自己合法注册并获得账号或 API Key"],
        ],
        [2700, 6660],
    )
    doc.add_heading("6. 安装主程序", level=2)
    for step in [
        "解压客户交付 ZIP，先打开“文档”目录阅读本手册。",
        "双击“AI 工具部署助手_1.0.0_aarch64.dmg”。",
        "把“AI 工具部署助手”拖入 Applications（应用程序）目录。",
        "从应用程序目录打开软件。不要从 DMG 窗口长期运行。",
        "如 macOS 提示来源或签名异常，先联系卖家核对安装包 SHA-256，不要关闭系统安全保护。",
    ]:
        numbered_paragraph(doc, step, decimal_num)

    doc.add_heading("7. 输入激活码", level=2)
    for step in [
        "启动后进入激活页。",
        "粘贴卖家单独发送的 ADA- 开头兑换码。",
        "点击“激活”，等待显示剩余授权天数。",
        "兑换码会绑定当前设备；换机前先联系卖家确认是否还有重置次数。",
    ]:
        numbered_paragraph(doc, step, decimal_num)

    doc.add_heading("第三部分：安装三款 AI 工具", level=1)
    doc.add_heading("8. 通用安装流程", level=2)
    for step in [
        "在“选择要部署的工具”区域找到目标工具。",
        "先点击“查看详情”，确认版本、安装类型、文件大小和来源页面。",
        "点击“开始安装”并等待进度到 100%。",
        "安装后点击“诊断”，确认系统兼容性、命令路径和 doctor 结果。",
        "需要初始化时点击“打开终端”，按官方向导继续。",
    ]:
        numbered_paragraph(doc, step, decimal_num)

    add_table(
        doc,
        ["工具", "首次启动", "主要用途", "推荐配置方式"],
        [
            ["Claude Code", "claude", "终端编程代理", "CC Switch 或 Anthropic 官方登录"],
            ["OpenClaw", "openclaw onboard", "个人 AI 助手和消息平台", "官方 onboard；多供应商用 CC Switch"],
            ["Hermes Agent", "hermes setup", "Nous Research 自主智能体", "官方 setup；多供应商用 CC Switch"],
        ],
        [1750, 2250, 2660, 2700],
    )

    doc.add_heading("第四部分：三款工具接入大模型", level=1)
    doc.add_heading("9. 先安装 CC Switch", level=2)
    for step in [
        "打开“配套工具/CC-Switch-v3.16.2-macOS.dmg”。",
        "把 CC Switch 拖入 Applications 目录后打开。",
        "CC Switch 同时支持 Claude Code、OpenClaw 和 Hermes，不需要安装三套切换软件。",
        "在 CC Switch 中添加供应商时，只填写供应商官方提供的 API Base URL、API Key 和模型名。",
        "切换配置后重启对应终端或 CLI 会话；不要把 API Key 发给卖家。",
    ]:
        numbered_paragraph(doc, step, decimal_num)
    add_callout(
        doc,
        "关于 CC Switch",
        "本交付包包含官方发布的 CC Switch 3.16.2 macOS Universal DMG，SHA-256 为 "
        "004a28e79f60f34840e32ed54c394c1ffdb55f25e90dfdf8cc5c3435fb987d8b。"
        "项目采用 MIT 许可证，许可证原文已放入交付包。",
        LIGHT_BLUE,
    )

    doc.add_heading("10. Claude Code 傻瓜式配置", level=2)
    for step in [
        "在部署助手中安装 Claude Code，并点击“诊断”。",
        "打开 CC Switch，选择 Claude Code 标签。",
        "点击新增供应商，填写 API Base URL、API Key 和模型；这些值由你的模型供应商提供。",
        "点击启用供应商。",
        "打开终端，进入项目目录，运行 claude。",
        "如使用 Anthropic 官方账号，可按 Claude Code 官方登录提示操作，不必使用第三方 API。",
    ]:
        numbered_paragraph(doc, step, decimal_num)

    doc.add_heading("11. OpenClaw 傻瓜式配置", level=2)
    for step in [
        "在部署助手中安装 OpenClaw，并点击“诊断”。",
        "首次使用运行 openclaw onboard，按提示选择模型提供商并填写凭据。",
        "如果以后需要频繁切换供应商，打开 CC Switch 的 OpenClaw 标签添加配置。",
        "启用新配置后重新启动 OpenClaw 终端或相关服务。",
        "模型服务和消息平台仍可能需要联网，离线安装包不等于离线模型。",
    ]:
        numbered_paragraph(doc, step, decimal_num)

    doc.add_heading("12. Hermes Agent 傻瓜式配置", level=2)
    for step in [
        "在部署助手中安装 Hermes Agent，并确认 hermes --version 可用。",
        "首次使用运行 hermes setup，完成依赖和模型提供商设置。",
        "需要统一管理时，在 CC Switch 的 Hermes 标签中添加 API 地址、密钥和默认模型。",
        "关闭旧 Hermes 会话并重新启动，再发送一个简单问题测试。",
        "Hermes 的 Python、Node、ripgrep、ffmpeg 等依赖应以诊断结果和官方文档为准。",
    ]:
        numbered_paragraph(doc, step, decimal_num)

    doc.add_heading("第五部分：诊断、卸载和更新", level=1)
    doc.add_heading("13. 诊断功能", level=2)
    add_table(
        doc,
        ["结果", "含义", "处理"],
        [
            ["绿色", "系统、命令路径或 doctor 检查正常", "继续初始化"],
            ["黄色", "doctor 超时或等待登录/网络", "打开终端完成登录后重试"],
            ["红色", "找不到命令或系统不兼容", "重新安装并把错误信息发给卖家"],
        ],
        [1400, 3900, 4060],
        header_fill=LIGHT_GREEN,
    )
    doc.add_heading("14. 卸载功能", level=2)
    for item in [
        "“卸载”只删除由部署助手管理的独立二进制或离线包，不删除客户项目、模型密钥或工具用户配置。",
        "由官方脚本或系统安装器安装的版本，会提示使用操作系统或官方方式卸载。",
        "卸载前会二次确认；不要在工具正在运行或正在更新时卸载。",
        "安全校验会拒绝删除部署助手数据目录之外的文件。",
    ]:
        numbered_paragraph(doc, item, bullet_num)
    doc.add_heading("15. 助手更新", level=2)
    for step in [
        "点击页面底部“检查更新”。",
        "发现新版本后核对版本说明，再点击安装更新。",
        "更新包经过签名校验；安装完成后应用会重启。",
        "更新失败时保留当前版本，记录错误信息并联系卖家。",
    ]:
        numbered_paragraph(doc, step, decimal_num)

    doc.add_heading("第六部分：常见问题与售后", level=1)
    add_table(
        doc,
        ["问题", "优先处理方法"],
        [
            ["按钮点击无反应", "确认软件版本为 1.0.0 正式 Release，关闭后重开，再截图错误提示"],
            ["诊断显示未找到命令", "重新安装对应工具，打开新终端后再诊断"],
            ["doctor 超时", "可能等待登录、授权或网络；在终端完成后重试"],
            ["激活码无效", "检查是否完整粘贴；卖家后台查看停用、过期或设备绑定状态"],
            ["客户换电脑", "卖家后台重置兑换码绑定，再在新电脑激活"],
            ["模型无法回答", "检查供应商余额、API 地址、Key、模型名和服务状态"],
            ["CC Switch 切换未生效", "关闭旧终端/CLI 会话后重新打开"],
            ["需要退款", "按商品页事先约定处理；不要要求客户发送 API Key 作为证明"],
        ],
        [3000, 6360],
    )

    doc.add_heading("卖家售后收集信息清单", level=2)
    for item in [
        "macOS 版本与芯片型号",
        "部署助手版本",
        "目标工具名称和版本",
        "诊断结果截图或错误文字",
        "问题发生前的操作步骤",
    ]:
        numbered_paragraph(doc, item, bullet_num)
    add_callout(
        doc,
        "不要收集",
        "不要让客户发送 API Key、OAuth Token、账号密码、完整设备指纹、项目私有代码或服务器密钥。"
        "如截图含密钥，先让客户打码。",
    )

    doc.add_heading("附录 A：交付包结构", level=1)
    add_table(
        doc,
        ["路径", "用途"],
        [
            ["AI 工具部署助手_1.0.0_aarch64.dmg", "macOS Apple 芯片主程序"],
            ["配套工具/CC-Switch-v3.16.2-macOS.dmg", "三款工具共用模型配置助手"],
            ["文档/AI工具部署助手-客户交付与使用手册.pdf", "客户安装与使用说明"],
            ["开源许可证/CC-Switch-MIT-LICENSE.txt", "CC Switch 再分发许可证"],
            ["SHA256SUMS.txt", "安装文件完整性校验"],
            ["交付说明.txt", "最短使用入口"],
        ],
        [4050, 5310],
    )

    doc.add_heading("附录 B：官方资料", level=1)
    for source in [
        "Claude Code：https://code.claude.com/docs/en/setup",
        "OpenClaw：https://docs.openclaw.ai/install",
        "Hermes Agent：https://hermes-agent.nousresearch.com/docs/",
        "CC Switch：https://github.com/farion1231/cc-switch",
        "CC Switch 3.16.2：https://github.com/farion1231/cc-switch/releases/tag/v3.16.2",
    ]:
        numbered_paragraph(doc, source, bullet_num)

    final_section = doc.sections[-1]
    if len(doc.sections) == 1:
        final_section.start_type = WD_SECTION.NEW_PAGE
    doc.core_properties.title = "AI 工具部署助手客户交付与使用手册"
    doc.core_properties.subject = "客户交付、激活、安装、模型接入与售后操作"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
